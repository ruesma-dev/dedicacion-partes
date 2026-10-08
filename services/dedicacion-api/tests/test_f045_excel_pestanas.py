# tests/test_f045_excel_pestanas.py
"""Tests offline de F-045: el Excel desglosado en «Obras» y «Postventa».

El libro lleva «Obras», «Postventa» y «Detalle» (la hoja de F-040, sin
cambios, la tercera) y no lleva Resumen (D1 y D2 de la MANUAL T6).

Dos niveles (specs/F-045-excel-obras-postventa/design.md §7):

  - Contenido (`grupos_pestana` de `infrastructure/excel/contenido.py`):
    qué líneas lleva el grupo de cada trabajador en cada pestaña y la línea
    agregada con la otra parte (R4-R12, R18).
  - Libro (`infrastructure/excel/exporter.py`): hojas, formato, combinadas,
    bandas y cursiva de la agregada, y «Detalle» igual que en F-040
    (R1-R3, R13-R17).

Los cinco trabajadores son los del prototipo (design §5), con nombres y
obras inventados. Sin red, sin BBDD y sin ficheros: el libro vive en memoria.
"""
from __future__ import annotations

import io
from decimal import Decimal

import pytest
from domain.models import CuadranteTrabajador, Periodo
from infrastructure.excel.contenido import (
    AGREGADA_EN_OBRAS,
    AGREGADA_EN_POSTVENTA,
    grupos_detalle,
    grupos_pestana,
)
from infrastructure.excel.exporter import OpenpyxlExcelExporter
from openpyxl import load_workbook

from tests.test_f040_excel import _fila, _linea, _rgb, _xml_hoja

_PREFIJO = "Postv-"
_OBRAS, _POSTVENTA = False, True


def _muestra() -> list[CuadranteTrabajador]:
    """ALFA obras y postventa (OK), BETA solo obras (FALTA 10 %), GAMMA solo
    postventa (OK), DELTA sin carga y EPSILON VAR + dos postventas (EXCESO)."""
    return [
        _fila(1, "ALFA PRUEBA UNO",
              _linea("0101", "Obra Ficticia Norte", "60"),
              _linea("0102", "Obra Ficticia Sur", "25"),
              _linea("0702", "Hotel Inventado", "15", postventa=True)),
        _fila(2, "BETA PRUEBA DOS",
              _linea("0101", "Obra Ficticia Norte", "50"),
              _linea("0103", "Obra Ficticia Este", "40")),
        _fila(3, "GAMMA PRUEBA TRES",
              _linea("0702", "Hotel Inventado", "70", postventa=True),
              _linea("0704", "Residencia Inventada", "30", postventa=True),
              categoria=None),
        _fila(4, "DELTA PRUEBA CUATRO"),
        _fila(5, "EPSILON PRUEBA CINCO",
              _linea("VAR-29", "Varios partida 29", "60"),
              _linea("0702", "Hotel Inventado", "30", postventa=True),
              _linea("0705", "Nave Inventada", "20", postventa=True)),
    ]


def _por_nombre(postventa: bool, filas=None, prefijo=_PREFIJO) -> dict:
    filas = _muestra() if filas is None else filas
    return {g.empleado: g for g in grupos_pestana(filas, prefijo, postventa)}


def _filas(grupo) -> list[tuple]:
    return [(ln.codigo, ln.obra, ln.porcentaje, ln.agregada)
            for ln in grupo.lineas]


# ----------------------------------------------------------------------
# Contenido
# ----------------------------------------------------------------------
def test_f045_r4_mismos_trabajadores_y_orden_en_las_dos_pestanas():
    for filas in (_muestra(), list(reversed(_muestra()))):
        esperado = [f.trabajador.nombre for f in filas]
        obras = grupos_pestana(filas, _PREFIJO, _OBRAS)
        postv = grupos_pestana(filas, _PREFIJO, _POSTVENTA)
        assert [g.empleado for g in obras] == esperado
        assert [g.empleado for g in postv] == esperado
        assert [g.categoria for g in obras] == [g.categoria for g in postv]


def test_f045_r5_obras_lleva_sus_lineas_en_orden_con_var_como_obra():
    obras = _por_nombre(_OBRAS)
    assert _filas(obras["BETA PRUEBA DOS"]) == [
        ("0101", "Obra Ficticia Norte", Decimal(50), False),
        ("0103", "Obra Ficticia Este", Decimal(40), False)]
    var = obras["EPSILON PRUEBA CINCO"].lineas[0]
    assert (var.codigo, var.obra, var.nombre, var.porcentaje, var.agregada) == (
        "VAR-29", "Varios partida 29", "Varios partida 29", Decimal(60), False)
    # Postventa intercalada: las obras conservan su orden relativo, sin reordenar.
    omega = _fila(9, "OMEGA",
                  _linea("0200", "B", "40"),
                  _linea("0100", "Pv", "20", postventa=True),
                  _linea("0150", "A", "40"))
    [g] = grupos_pestana([omega], _PREFIJO, _OBRAS)
    assert [ln.codigo for ln in g.lineas] == ["0200", "0150", "POSTVENTA"]


def test_f045_r6_postventa_lleva_sus_lineas_en_orden_con_prefijo():
    postv = _por_nombre(_POSTVENTA)
    gamma = postv["GAMMA PRUEBA TRES"]
    assert _filas(gamma) == [
        ("Postv-0702", "Postv-Hotel Inventado", Decimal(70), False),
        ("Postv-0704", "Postv-Residencia Inventada", Decimal(30), False)]
    assert gamma.lineas[0].nombre == "Hotel Inventado"
    assert [ln.codigo for ln in postv["EPSILON PRUEBA CINCO"].lineas] == [
        "Postv-0702", "Postv-0705", "OBRAS"]


def test_f045_r7_agregada_unica_al_final_con_rotulos_y_suma():
    assert AGREGADA_EN_OBRAS == ("POSTVENTA", "RESTO POSTVENTA")
    assert AGREGADA_EN_POSTVENTA == ("OBRAS", "RESTO OBRAS")
    obras, postv = _por_nombre(_OBRAS), _por_nombre(_POSTVENTA)
    alfa_o, alfa_p = obras["ALFA PRUEBA UNO"], postv["ALFA PRUEBA UNO"]
    assert _filas(alfa_o) == [
        ("0101", "Obra Ficticia Norte", Decimal(60), False),
        ("0102", "Obra Ficticia Sur", Decimal(25), False),
        ("POSTVENTA", "RESTO POSTVENTA", Decimal(15), True)]
    assert _filas(alfa_p) == [
        ("Postv-0702", "Postv-Hotel Inventado", Decimal(15), False),
        ("OBRAS", "RESTO OBRAS", Decimal(85), True)]
    assert alfa_o.lineas[-1].nombre == "" and alfa_p.lineas[-1].nombre == ""
    # La suma es de TODAS las líneas de la otra parte.
    eps_o = obras["EPSILON PRUEBA CINCO"]
    assert _filas(eps_o)[-1] == ("POSTVENTA", "RESTO POSTVENTA", Decimal(50), True)
    assert _filas(postv["EPSILON PRUEBA CINCO"])[-1] == (
        "OBRAS", "RESTO OBRAS", Decimal(60), True)
    # Una sola agregada por grupo, y siempre la última.
    for grupos in (obras, postv):
        for g in grupos.values():
            marcas = [ln.agregada for ln in g.lineas]
            assert sum(marcas) <= 1
            assert True not in marcas[:-1]


def test_f045_r7_suma_exacta_en_decimal():
    fila = _fila(6, "THETA",
                 _linea("0101", "A", "33.33"),
                 _linea("0702", "B", "10.1", postventa=True),
                 _linea("0703", "C", "20.2", postventa=True),
                 _linea("0102", "D", "36.37"))
    [obras] = grupos_pestana([fila], _PREFIJO, _OBRAS)
    [postv] = grupos_pestana([fila], _PREFIJO, _POSTVENTA)
    resto_pv, resto_ob = obras.lineas[-1], postv.lineas[-1]
    assert isinstance(resto_pv.porcentaje, Decimal)
    assert resto_pv.porcentaje == Decimal("30.3")      # no 30.299999…
    assert isinstance(resto_ob.porcentaje, Decimal)
    assert resto_ob.porcentaje == Decimal("69.70")


def test_f045_r8_sin_la_otra_parte_no_hay_agregada():
    obras, postv = _por_nombre(_OBRAS), _por_nombre(_POSTVENTA)
    beta, gamma = obras["BETA PRUEBA DOS"], postv["GAMMA PRUEBA TRES"]
    for g in (beta, gamma):
        assert not any(ln.agregada for ln in g.lineas)
        assert len(g.lineas) == 2
    assert "POSTVENTA" not in [ln.codigo for ln in beta.lineas]
    assert "OBRAS" not in [ln.codigo for ln in gamma.lineas]


def test_f045_r9_solo_la_otra_parte_una_fila_agregada():
    obras, postv = _por_nombre(_OBRAS), _por_nombre(_POSTVENTA)
    assert _filas(obras["GAMMA PRUEBA TRES"]) == [
        ("POSTVENTA", "RESTO POSTVENTA", Decimal(100), True)]
    assert _filas(postv["BETA PRUEBA DOS"]) == [
        ("OBRAS", "RESTO OBRAS", Decimal(90), True)]


def test_f045_r10_sin_carga_una_fila_vacia_en_cada_pestana():
    for postventa in (_OBRAS, _POSTVENTA):
        delta = _por_nombre(postventa)["DELTA PRUEBA CUATRO"]
        assert _filas(delta) == [("", "", None, False)]
        assert delta.lineas[0].nombre == ""
        assert delta.total == Decimal(0)
        assert delta.desviacion is None
        assert delta.estado == "SIN CARGA"


def test_f045_r11_total_desviacion_y_estado_del_trabajador_completo():
    completos = grupos_detalle(_muestra(), _PREFIJO)
    for postventa in (_OBRAS, _POSTVENTA):
        grupos = grupos_pestana(_muestra(), _PREFIJO, postventa)
        assert [(g.empleado, g.categoria, g.total, g.desviacion, g.estado)
                for g in grupos] == [
            (g.empleado, g.categoria, g.total, g.desviacion, g.estado)
            for g in completos]
    assert [g.estado for g in completos] == [
        "OK", "FALTA 10%", "OK", "SIN CARGA", "EXCESO 10%"]
    assert completos[2].categoria == ""


@pytest.mark.parametrize("postventa", [_OBRAS, _POSTVENTA])
def test_f045_r12_la_columna_pct_del_grupo_suma_el_total(postventa):
    theta = _fila(6, "THETA",
                  _linea("0101", "A", "33.33"),
                  _linea("0702", "B", "10.1", postventa=True),
                  _linea("0703", "C", "20.2", postventa=True),
                  _linea("0102", "D", "36.36"))
    for g in grupos_pestana([*_muestra(), theta], _PREFIJO, postventa):
        pcts = [ln.porcentaje for ln in g.lineas if ln.porcentaje is not None]
        if g.estado == "SIN CARGA":
            assert pcts == []
            continue
        assert sum(pcts, Decimal(0)) == g.total, g.empleado


def test_f045_r12_la_suma_coincide_en_obras_y_en_postventa():
    obras, postv = _por_nombre(_OBRAS), _por_nombre(_POSTVENTA)
    for nombre in obras:
        suma_o = sum((ln.porcentaje for ln in obras[nombre].lineas
                      if ln.porcentaje is not None), Decimal(0))
        suma_p = sum((ln.porcentaje for ln in postv[nombre].lineas
                      if ln.porcentaje is not None), Decimal(0))
        assert suma_o == suma_p, nombre


def test_f045_r5_d5_var_en_obras_y_sumado_en_resto_obras():
    obras, postv = _por_nombre(_OBRAS), _por_nombre(_POSTVENTA)
    assert "VAR-29" in [ln.codigo for ln in obras["EPSILON PRUEBA CINCO"].lineas]
    codigos_pv = [ln.codigo for g in postv.values() for ln in g.lineas]
    assert not any("VAR" in c for c in codigos_pv)
    # RESTO OBRAS de EPSILON es exactamente su VAR-29 (60).
    assert postv["EPSILON PRUEBA CINCO"].lineas[-1].porcentaje == Decimal(60)


def test_f045_r18_el_prefijo_sale_de_la_configuracion():
    postv = _por_nombre(_POSTVENTA, prefijo="PV_")
    assert _filas(postv["ALFA PRUEBA UNO"]) == [
        ("PV_0702", "PV_Hotel Inventado", Decimal(15), False),
        ("OBRAS", "RESTO OBRAS", Decimal(85), True)]
    obras = _por_nombre(_OBRAS, prefijo="PV_")
    assert _filas(obras["ALFA PRUEBA UNO"])[-1] == (
        "POSTVENTA", "RESTO POSTVENTA", Decimal(15), True)


# ----------------------------------------------------------------------
# Libro
# ----------------------------------------------------------------------
_PERIODO = Periodo(anio=2026, mes=9)
_BLANCO, _AZUL = "FFFFFFFF", "FFDDEBF7"
_CABECERA = ["Empleado", "Categoría", "Código", "Obra", "% dedicación",
             "Total empleado", "Desviación", "Estado"]

#: (primera, última) fila de cada grupo de `_muestra()`, por hoja.
_GRUPOS = {
    "Obras": [(3, 5), (6, 7), (8, 8), (9, 9), (10, 11)],
    "Postventa": [(3, 4), (5, 5), (6, 7), (8, 8), (9, 11)],
}
#: Filas de la línea agregada, por hoja.
_AGREGADAS = {"Obras": {5, 8, 11}, "Postventa": {4, 5, 11}}
#: Número de `sheet<n>.xml` de cada hoja.
_N_HOJA = {"Obras": 1, "Postventa": 2, "Detalle": 3}
#: (primera, última) fila de cada grupo de `_muestra()` en «Detalle».
_GRUPOS_DETALLE = [(3, 5), (6, 7), (8, 9), (10, 10), (11, 13)]


def _exportar(filas=None, prefijo=_PREFIJO, periodo=_PERIODO) -> bytes:
    filas = _muestra() if filas is None else filas
    return OpenpyxlExcelExporter(prefijo_postventa=prefijo).exportar(periodo, filas)


def _libro(contenido: bytes):
    return load_workbook(io.BytesIO(contenido))


def _columna(hoja, col: int) -> list:
    return [hoja.cell(r, col).value for r in range(3, hoja.max_row + 1)]


def _empleados(contenido: bytes, nombre: str) -> list:
    """Columna A desde el XML: `load_workbook` vacía las no ancla (F-040 §5)."""
    celdas, _ = _xml_hoja(contenido, _N_HOJA[nombre])
    return [celdas[f"A{r}"]["v"] for r in range(3, 12)]


def test_f045_r1_tres_hojas_obras_postventa_y_detalle():
    assert _libro(_exportar()).sheetnames == ["Obras", "Postventa", "Detalle"]
    assert _libro(_exportar([])).sheetnames == ["Obras", "Postventa", "Detalle"]


def test_f045_r2_titulos_cabecera_en_2_y_datos_desde_3():
    libro = _libro(_exportar())
    assert libro["Obras"]["A1"].value == "OBRAS · Septiembre 2026"
    assert libro["Postventa"]["A1"].value == "POSTVENTA · Septiembre 2026"
    assert libro["Detalle"]["A1"].value == (
        "DETALLE DE DEDICACIÓN · Septiembre 2026")
    for hoja in libro.worksheets:
        assert hoja["A1"].font.bold and hoja["A1"].font.size == 13
        assert hoja["A2"].value == "Empleado"
        assert hoja["A3"].value == "ALFA PRUEBA UNO"
    enero = _libro(_exportar([], periodo=Periodo(2027, 1)))
    assert enero["Obras"]["A1"].value == "OBRAS · Enero 2027"
    assert enero["Postventa"]["A1"].value == "POSTVENTA · Enero 2027"
    assert enero["Detalle"]["A1"].value == "DETALLE DE DEDICACIÓN · Enero 2027"


@pytest.mark.parametrize("nombre", ["Obras", "Postventa"])
def test_f045_r3_cabecera_autofiltro_paneles_anchos_e_impresion(nombre):
    hoja = _libro(_exportar())[nombre]
    assert [c.value for c in hoja[2]] == _CABECERA
    for celda in hoja[2]:
        assert _rgb(celda.fill.fgColor.rgb) == "FF1F3864"
        assert celda.font.bold and _rgb(celda.font.color.rgb) == "FFFFFFFF"
    assert hoja.max_row == 11
    assert hoja.auto_filter.ref == "A2:H11"
    assert hoja.freeze_panes == "A3"
    assert tuple(hoja.column_dimensions[c].width for c in "ABCDEFGH") == (
        34, 22, 12, 44, 13, 15, 12, 16)
    assert hoja.page_setup.orientation == "landscape"
    assert hoja.sheet_properties.pageSetUpPr.fitToPage is True
    assert hoja.page_setup.fitToWidth == 1
    assert hoja.page_setup.fitToHeight == 0
    assert hoja.print_title_rows == "$1:$2"
    assert _libro(_exportar([]))[nombre].auto_filter.ref == "A2:H2"


def test_f045_r4_r10_filas_de_la_pestana_obras():
    contenido = _exportar()
    obras = _libro(contenido)["Obras"]
    assert _empleados(contenido, "Obras") == (
        ["ALFA PRUEBA UNO"] * 3 + ["BETA PRUEBA DOS"] * 2
        + ["GAMMA PRUEBA TRES", "DELTA PRUEBA CUATRO"]
        + ["EPSILON PRUEBA CINCO"] * 2)
    assert _columna(obras, 3) == [
        "0101", "0102", "POSTVENTA", "0101", "0103", "POSTVENTA", None,
        "VAR-29", "POSTVENTA"]
    assert _columna(obras, 4) == [
        "Obra Ficticia Norte", "Obra Ficticia Sur", "RESTO POSTVENTA",
        "Obra Ficticia Norte", "Obra Ficticia Este", "RESTO POSTVENTA", None,
        "Varios partida 29", "RESTO POSTVENTA"]
    assert _columna(obras, 5) == [0.6, 0.25, 0.15, 0.5, 0.4, 1, None, 0.6, 0.5]
    assert [c.value for c in obras[9]][2:] == [None, None, None, 0, None,
                                               "SIN CARGA"]


def test_f045_r4_r10_filas_de_la_pestana_postventa():
    contenido = _exportar()
    postv = _libro(contenido)["Postventa"]
    assert _empleados(contenido, "Postventa") == (
        ["ALFA PRUEBA UNO"] * 2 + ["BETA PRUEBA DOS"]
        + ["GAMMA PRUEBA TRES"] * 2 + ["DELTA PRUEBA CUATRO"]
        + ["EPSILON PRUEBA CINCO"] * 3)
    assert _columna(postv, 3) == [
        "Postv-0702", "OBRAS", "OBRAS", "Postv-0702", "Postv-0704", None,
        "Postv-0702", "Postv-0705", "OBRAS"]
    assert _columna(postv, 4) == [
        "Postv-Hotel Inventado", "RESTO OBRAS", "RESTO OBRAS",
        "Postv-Hotel Inventado", "Postv-Residencia Inventada", None,
        "Postv-Hotel Inventado", "Postv-Nave Inventada", "RESTO OBRAS"]
    assert _columna(postv, 5) == [0.15, 0.85, 0.9, 0.7, 0.3, None, 0.3, 0.2, 0.6]
    assert [c.value for c in postv[8]][2:] == [None, None, None, 0, None,
                                               "SIN CARGA"]


def test_f045_r11_total_desviacion_y_estado_iguales_en_las_tres_hojas():
    libro = _libro(_exportar())
    det = libro["Detalle"]
    for nombre, grupos in _GRUPOS.items():
        hoja = libro[nombre]
        for (ini, _fin), (ini_det, _) in zip(grupos, _GRUPOS_DETALLE, strict=True):
            for col in (1, 6, 7, 8):
                assert hoja.cell(ini, col).value == det.cell(ini_det, col).value, (
                    nombre, ini, col)
        assert [hoja.cell(ini, 7).value for ini, _ in grupos] == [
            0, -0.1, 0, None, pytest.approx(0.1)], nombre


def test_f045_r12_la_columna_pct_de_cada_grupo_suma_su_total_en_el_libro():
    libro = _libro(_exportar())
    for nombre, grupos in _GRUPOS.items():
        hoja = libro[nombre]
        for ini, fin in grupos:
            pcts = [hoja.cell(r, 5).value or 0 for r in range(ini, fin + 1)]
            assert sum(pcts) == pytest.approx(hoja.cell(ini, 6).value), (
                nombre, ini)


def test_f045_r13_valor_en_todas_las_filas_y_combinadas_con_la_agregada():
    contenido = _exportar()
    for nombre, grupos in _GRUPOS.items():
        celdas, rangos = _xml_hoja(contenido, _N_HOJA[nombre])
        for ini, fin in grupos:
            for col in "ABFGH":
                valores = {celdas[f"{col}{r}"]["v"] for r in range(ini, fin + 1)}
                assert len(valores) == 1, (nombre, col, ini, valores)
                # Empleado, Total y Estado nunca vacíos, tampoco en la agregada.
                assert None not in valores or col in "BG", (nombre, col, ini)
        esperados = {f"{c}{ini}:{c}{fin}" for ini, fin in grupos if fin > ini
                     for c in "ABFGH"}
        assert set(rangos) == esperados and len(rangos) == len(esperados), nombre
        hoja = _libro(contenido)[nombre]
        for ini, fin in grupos:
            if fin > ini:
                for c in "ABFGH":
                    assert hoja[f"{c}{ini}"].alignment.vertical == "center"
    # La agregada es una fila más del grupo: A3:A5 en Obras, A3:A4 en Postventa.
    assert "A3:A5" in _xml_hoja(contenido, 1)[1]
    assert "A3:A4" in _xml_hoja(contenido, 2)[1]
    # «Detalle», como en F-040: las combinadas de sus grupos, sin agregada.
    _, rangos_det = _xml_hoja(contenido, _N_HOJA["Detalle"])
    esperados_det = {f"{c}{ini}:{c}{fin}" for ini, fin in _GRUPOS_DETALLE
                     if fin > ini for c in "ABFGH"}
    assert set(rangos_det) == esperados_det
    assert len(rangos_det) == len(esperados_det)


def test_f045_r14_bandas_por_hoja_y_linea_gruesa_bajo_cada_trabajador():
    contenido = _exportar()
    for nombre, grupos in _GRUPOS.items():
        celdas, _ = _xml_hoja(contenido, _N_HOJA[nombre])
        for n, (ini, fin) in enumerate(grupos):
            banda = (_BLANCO, _AZUL)[n % 2]
            for r in range(ini, fin + 1):
                for col in "ABCDEFGH":
                    assert celdas[f"{col}{r}"]["fill"] == banda, (nombre, col, r)
            for col in "ABCDEFGH":
                assert celdas[f"{col}{fin}"]["bottom"] == (
                    "medium", "FF000000"), (nombre, col)
            for r in range(ini, fin):
                for col in "CDE":
                    assert celdas[f"{col}{r}"]["bottom"] == ("thin", "FFBFBFBF")


def test_f045_r15_cursiva_solo_en_codigo_obra_y_pct_de_la_agregada():
    libro = _libro(_exportar())
    for hoja in libro.worksheets:
        agregadas = _AGREGADAS.get(hoja.title, set())
        for fila in hoja.iter_rows():
            for celda in fila:
                esperada = celda.row in agregadas and celda.column in (3, 4, 5)
                assert bool(celda.font.italic) is esperada, (
                    hoja.title, celda.coordinate)
                if esperada:
                    assert not celda.font.bold
    assert libro["Obras"]["D5"].value == "RESTO POSTVENTA"
    assert libro["Postventa"]["D4"].value == "RESTO OBRAS"


def test_f045_r16_pct_de_la_agregada_fraccion_formato_y_sin_formulas():
    libro = _libro(_exportar())
    obras, postv = libro["Obras"], libro["Postventa"]
    assert (obras["E5"].value, obras["E5"].number_format) == (0.15, "0%")
    assert (postv["E4"].value, postv["E4"].number_format) == (0.85, "0%")
    theta = _fila(6, "THETA",
                  _linea("0101", "A", "33.33"),
                  _linea("0702", "B", "10.1", postventa=True),
                  _linea("0703", "C", "20.2", postventa=True),
                  _linea("0102", "D", "36.37"))
    libro2 = _libro(_exportar([theta]))
    o2, p2 = libro2["Obras"], libro2["Postventa"]
    assert o2["D5"].value == "RESTO POSTVENTA"
    assert (o2["E5"].value, o2["E5"].number_format) == (
        pytest.approx(0.303), "0.00%")
    assert p2["D5"].value == "RESTO OBRAS"
    assert (p2["E5"].value, p2["E5"].number_format) == (
        pytest.approx(0.697), "0.00%")
    for lib in (libro, libro2):
        for hoja in lib.worksheets:
            for fila in hoja.iter_rows():
                for celda in fila:
                    assert celda.data_type != "f", (hoja.title, celda.coordinate)
                    assert not str(celda.value or "").startswith("=")


def test_f045_r17_detalle_igual_que_en_f040():
    contenido = _exportar()
    det = _libro(contenido)["Detalle"]
    grupos = grupos_detalle(_muestra(), _PREFIJO)
    assert det["A1"].value == "DETALLE DE DEDICACIÓN · Septiembre 2026"
    assert [c.value for c in det[2]] == _CABECERA
    assert det.max_row == 13
    assert det.auto_filter.ref == "A2:H13"
    lineas = [ln for g in grupos for ln in g.lineas]
    assert len(lineas) == 11
    for r, ln in zip(range(3, 14), lineas, strict=True):
        assert (det.cell(r, 3).value or "") == ln.codigo, r
        assert (det.cell(r, 4).value or "") == ln.obra, r
        esperado = None if ln.porcentaje is None else float(ln.porcentaje) / 100
        assert det.cell(r, 5).value == pytest.approx(esperado), r
    # Postventa intercalada con su prefijo, y ninguna línea agregada.
    assert [det.cell(r, 3).value for r in (3, 4, 5)] == [
        "0101", "0102", "Postv-0702"]
    assert det["D5"].value == "Postv-Hotel Inventado"
    codigos = [det.cell(r, 3).value for r in range(3, 14)]
    obras = [det.cell(r, 4).value or "" for r in range(3, 14)]
    assert "POSTVENTA" not in codigos and "OBRAS" not in codigos
    assert not any(o.startswith("RESTO") for o in obras)
    celdas, _ = _xml_hoja(contenido, _N_HOJA["Detalle"])
    for n, ((ini, fin), g) in enumerate(zip(_GRUPOS_DETALLE, grupos, strict=True)):
        assert det.cell(ini, 1).value == g.empleado
        assert det.cell(ini, 6).value == pytest.approx(float(g.total) / 100)
        assert det.cell(ini, 8).value == g.estado
        banda = (_BLANCO, _AZUL)[n % 2]
        for r in range(ini, fin + 1):
            for col in "ABCDEFGH":
                assert celdas[f"{col}{r}"]["fill"] == banda, (col, r)
        for col in "ABCDEFGH":
            assert celdas[f"{col}{fin}"]["bottom"] == ("medium", "FF000000"), col


def test_f045_r18_el_prefijo_de_la_configuracion_llega_al_libro():
    libro = _libro(_exportar(prefijo="PV_"))
    postv = libro["Postventa"]
    assert (postv["C3"].value, postv["D3"].value) == (
        "PV_0702", "PV_Hotel Inventado")
    assert (postv["C4"].value, postv["D4"].value) == ("OBRAS", "RESTO OBRAS")
    assert libro["Obras"]["C5"].value == "POSTVENTA"
    assert libro["Detalle"]["C5"].value == "PV_0702"
