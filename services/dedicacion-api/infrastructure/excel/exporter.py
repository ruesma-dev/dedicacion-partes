# infrastructure/excel/exporter.py
"""Exportación a Excel del periodo.

Genera dos hojas compatibles con los consumidores de la plantilla actual:
  - "Detalle": salida normalizada (una fila por trabajador y obra), con
    las mismas columnas que la hoja Detalle del Excel v14. Las líneas de
    postventa se exportan con el convenio "Postv-<código>".
  - "Resumen": una fila por trabajador con total y estado.
"""
from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from domain.estados import calcular_desviacion, calcular_estado
from domain.models import CuadranteTrabajador, EstadoTrabajador, Periodo

_MESES = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]

_CABECERA_FILL = PatternFill("solid", fgColor="1F3864")
_CABECERA_FONT = Font(color="FFFFFF", bold=True)

_ESTADO_TEXTO = {
    EstadoTrabajador.OK: "OK",
    EstadoTrabajador.SIN_CARGA: "SIN CARGA",
}


class OpenpyxlExcelExporter:
    def __init__(self, prefijo_postventa: str = "Postv-") -> None:
        self._prefijo = prefijo_postventa

    # ------------------------------------------------------------------
    def exportar(
        self, periodo: Periodo, filas: list[CuadranteTrabajador]
    ) -> bytes:
        libro = Workbook()
        self._hoja_detalle(libro.active, periodo, filas)
        self._hoja_resumen(libro.create_sheet("Resumen"), periodo, filas)
        buffer = io.BytesIO()
        libro.save(buffer)
        return buffer.getvalue()

    # ------------------------------------------------------------------
    def _hoja_detalle(
        self, hoja, periodo: Periodo, filas: list[CuadranteTrabajador]
    ) -> None:
        hoja.title = "Detalle"
        hoja.append(
            [f"DETALLE DE DEDICACIÓN · {_MESES[periodo.mes - 1]} {periodo.anio}"]
        )
        hoja["A1"].font = Font(bold=True, size=13)
        hoja.append([])
        cabecera = [
            "Empleado", "Categoría", "Código", "Obra", "Obra(código)",
            "% dedicación", "Total empleado", "Desviación", "Estado",
        ]
        hoja.append(cabecera)
        for celda in hoja[3]:
            celda.fill = _CABECERA_FILL
            celda.font = _CABECERA_FONT
            celda.alignment = Alignment(vertical="center")

        for fila in filas:
            total = fila.total
            estado = calcular_estado(total, len(fila.lineas))
            desviacion = calcular_desviacion(total, len(fila.lineas))
            texto_estado = self._texto_estado(estado, desviacion)
            for linea in fila.lineas:
                codigo = (
                    f"{self._prefijo}{linea.cod}" if linea.es_postventa else linea.cod
                )
                nombre_obra = (
                    f"{self._prefijo}{linea.descripcion}"
                    if linea.es_postventa
                    else linea.descripcion
                )
                hoja.append(
                    [
                        fila.trabajador.nombre,
                        fila.trabajador.categoria or "",
                        codigo,
                        nombre_obra,
                        f"{nombre_obra} ({codigo})",
                        _pct(linea.porcentaje),
                        _pct(total),
                        _pct(desviacion),
                        texto_estado,
                    ]
                )
            if not fila.lineas:
                hoja.append(
                    [
                        fila.trabajador.nombre,
                        fila.trabajador.categoria or "",
                        "", "", "",
                        None,
                        _pct(Decimal("0")),
                        None,
                        self._texto_estado(estado, desviacion),
                    ]
                )
        self._formatear(hoja, num_cols=9, cols_pct=(6, 7, 8))

    def _hoja_resumen(
        self, hoja, periodo: Periodo, filas: list[CuadranteTrabajador]
    ) -> None:
        hoja.append(
            [f"RESUMEN · {_MESES[periodo.mes - 1]} {periodo.anio}"]
        )
        hoja["A1"].font = Font(bold=True, size=13)
        hoja.append([])
        hoja.append(["Empleado", "Categoría", "Obras", "Total %", "Estado"])
        for celda in hoja[3]:
            celda.fill = _CABECERA_FILL
            celda.font = _CABECERA_FONT
        for fila in filas:
            total = fila.total
            estado = calcular_estado(total, len(fila.lineas))
            desviacion = calcular_desviacion(total, len(fila.lineas))
            partes = []
            for linea in fila.lineas:
                codigo = (
                    f"{self._prefijo}{linea.cod}" if linea.es_postventa else linea.cod
                )
                partes.append(f"{codigo} = {linea.porcentaje.normalize()}%")
            hoja.append(
                [
                    fila.trabajador.nombre,
                    fila.trabajador.categoria or "",
                    " + ".join(partes),
                    _pct(total),
                    self._texto_estado(estado, desviacion),
                ]
            )
        self._formatear(hoja, num_cols=5, cols_pct=(4,))

    # ------------------------------------------------------------------
    def _texto_estado(
        self, estado: EstadoTrabajador, desviacion: Decimal
    ) -> str:
        if estado in _ESTADO_TEXTO:
            return _ESTADO_TEXTO[estado]
        magnitud = abs(desviacion).normalize()
        if estado is EstadoTrabajador.FALTA:
            return f"FALTA {magnitud}%"
        return f"EXCESO {magnitud}%"

    @staticmethod
    def _formatear(hoja, num_cols: int, cols_pct: tuple[int, ...]) -> None:
        anchos = {1: 38, 2: 24, 3: 14, 4: 38, 5: 46}
        for col in range(1, num_cols + 1):
            letra = get_column_letter(col)
            hoja.column_dimensions[letra].width = anchos.get(col, 14)
        for fila in hoja.iter_rows(min_row=4):
            for idx in cols_pct:
                celda = fila[idx - 1]
                if celda.value is not None:
                    celda.number_format = "0.00%"
        hoja.freeze_panes = "A4"


def _pct(valor: Decimal) -> float:
    """Excel espera la fracción (0.25) para el formato 0.00%."""
    return float(valor) / 100.0
